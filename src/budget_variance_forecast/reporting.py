"""生成完全脱敏的 Excel 输入文件和管理报告。"""

from __future__ import annotations

from dataclasses import dataclass
from numbers import Real
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .forecast import ForecastResult, build_rolling_forecast
from .synthetic import SyntheticDataset, build_synthetic_dataset
from .variance import VarianceResult, calculate_variances


REPORT_SHEETS = [
    "管理摘要",
    "利润差异桥",
    "月度趋势",
    "滚动预测",
    "预测明细",
    "差异明细",
    "数据质量",
]
INPUT_SHEETS = ["演示说明", "预算", "实际", "预测假设"]

NAVY = "17365D"
BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_RED = "FCE4D6"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"
BLACK = "222222"
GREEN = "008000"
THIN_GRAY = Side(style="thin", color="D9E1F2")


@dataclass(frozen=True)
class DemoArtifacts:
    """一键演示生成的两个 Excel 文件及核心结果。"""

    input_path: Path
    report_path: Path
    dataset: SyntheticDataset
    variance: VarianceResult
    forecast: ForecastResult


def generate_demo_artifacts(output_dir: str | Path = "demo_output") -> DemoArtifacts:
    """生成合成输入和完整管理报告，不读取任何外部业务数据。"""

    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    dataset = build_synthetic_dataset(actual_months=6)
    variance = calculate_variances(dataset.budget, dataset.actual)
    forecast = build_rolling_forecast(
        dataset.budget, dataset.actual, dataset.assumptions
    )
    input_path = directory / "合成预算与实际输入.xlsx"
    report_path = directory / "预算差异与滚动预测管理报告.xlsx"
    write_input_workbook(input_path, dataset)
    write_management_report(report_path, dataset, variance, forecast)
    return DemoArtifacts(input_path, report_path, dataset, variance, forecast)


def write_input_workbook(path: str | Path, dataset: SyntheticDataset) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)

    notice = workbook.create_sheet("演示说明")
    notice.append(["说明"])
    for message in (
        "本文件全部为程序生成的合成数据，不对应任何真实公司或业务记录。",
        "预算覆盖2026年12个月；实际数据截至2026年6月。",
        "金额单位为人民币元，数量为虚构业务单位。",
        "公开仓库不得加入真实预算、客户、产品、价格或财务记录。",
    ):
        notice.append([message])

    _write_dataframe(
        workbook.create_sheet("预算"), _budget_input_export(dataset.budget), "BudgetTable"
    )
    _write_dataframe(
        workbook.create_sheet("实际"), _actual_input_export(dataset.actual), "ActualTable"
    )
    _write_dataframe(
        workbook.create_sheet("预测假设"),
        _assumption_input_export(dataset.assumptions),
        "AssumptionTable",
    )
    _finish_workbook(workbook)
    workbook.save(path)
    return Path(path)


def write_management_report(
    path: str | Path,
    dataset: SyntheticDataset,
    variance: VarianceResult,
    forecast: ForecastResult,
) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {name: workbook.create_sheet(name) for name in REPORT_SHEETS}

    target = _annual_budget_operating_profit(dataset.budget)
    _build_executive_summary(sheets["管理摘要"], variance, forecast, target)
    _build_profit_bridge(sheets["利润差异桥"], variance)
    _build_monthly_trend(sheets["月度趋势"], variance)
    _build_forecast_summary(sheets["滚动预测"], forecast, target)
    _write_dataframe(sheets["预测明细"], _forecast_export(forecast.detail), "ForecastDetail")
    _write_dataframe(sheets["差异明细"], _variance_export(variance.detail), "VarianceDetail")
    _build_data_quality(sheets["数据质量"], dataset, variance, forecast)

    _finish_workbook(workbook)
    workbook.save(path)
    return Path(path)


def _build_executive_summary(
    sheet, variance: VarianceResult, forecast: ForecastResult, annual_target: float
) -> None:
    _title(sheet, "预算执行与滚动预测管理摘要", f"实际截至 {forecast.actual_through:%Y-%m}｜金额单位：人民币元")
    amounts = dict(zip(variance.summary["metric"], variance.summary["amount"]))
    base = forecast.summary.set_index("scenario").loc["基准"]
    cards = [
        ("上半年实际收入", amounts["实际收入"]),
        ("收入预算差异", amounts["收入差异"]),
        ("上半年实际经营利润", amounts["实际经营利润"]),
        ("经营利润预算差异", amounts["经营利润差异"]),
        ("基准情景全年收入", float(base["revenue"])),
        ("基准情景全年经营利润", float(base["operating_profit"])),
        ("年度经营利润目标", annual_target),
        ("基准情景目标差额", float(base["operating_profit"] - annual_target)),
    ]
    sheet.append([])
    sheet.append(["指标", "金额", "判断"])
    for label, amount in cards:
        judgment = "有利" if "差异" in label and amount >= 0 else "不利" if "差异" in label else ""
        sheet.append([label, round(amount, 2), judgment])
    _style_table_area(sheet, 5, 5 + len(cards), 3)
    sheet["B6"].number_format = '¥#,##0;[Red](¥#,##0);-'
    for row in range(6, 6 + len(cards)):
        sheet.cell(row, 2).number_format = '¥#,##0;[Red](¥#,##0);-'
    sheet["E5"] = "管理层关注"
    sheet["E5"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["E5"].font = Font(name="Microsoft YaHei", bold=True, color=WHITE)
    notes = [
        "收入高于预算，但经营利润低于预算。",
        "主要不利项来自单位变动成本和固定费用。",
        "乐观与谨慎情景只改变未来月份，历史实际保持不变。",
        "情景结果是条件模拟，不代表因果预测或业绩承诺。",
    ]
    for index, note in enumerate(notes, start=6):
        sheet.cell(index, 5, note)
    sheet.column_dimensions["A"].width = 27
    sheet.column_dimensions["B"].width = 19
    sheet.column_dimensions["C"].width = 12
    sheet.column_dimensions["D"].width = 3
    sheet.column_dimensions["E"].width = 55


def _build_profit_bridge(sheet, variance: VarianceResult) -> None:
    _title(sheet, "上半年经营利润差异桥", "正数为有利影响，负数为不利影响")
    amounts = dict(zip(variance.summary["metric"], variance.summary["amount"]))
    rows = [
        ["预算经营利润", amounts["预算经营利润"]],
        ["销量影响", amounts["销量对利润影响"]],
        ["产品结构影响", amounts["产品结构对利润影响"]],
        ["价格影响", amounts["价格对利润影响"]],
        ["单位成本影响", amounts["单位成本对利润影响"]],
        ["固定费用影响", amounts["固定费用对利润影响"]],
        ["实际经营利润", amounts["实际经营利润"]],
        ["勾稽差异", amounts["利润桥勾稽差异"]],
    ]
    sheet.append([])
    sheet.append(["利润桥项目", "金额"])
    for label, amount in rows:
        sheet.append([label, round(amount, 2)])
    _style_table_area(sheet, 5, 5 + len(rows), 2)
    for row in range(6, 6 + len(rows)):
        sheet.cell(row, 2).number_format = '¥#,##0;[Red](¥#,##0);-'
    chart = BarChart()
    chart.title = "经营利润差异构成"
    chart.y_axis.title = "人民币元"
    chart.add_data(Reference(sheet, min_col=2, min_row=5, max_row=12), titles_from_data=True)
    chart.series[0].tx = SeriesLabel(v="金额")
    chart.set_categories(Reference(sheet, min_col=1, min_row=6, max_row=12))
    chart.height = 8
    chart.width = 16
    sheet.add_chart(chart, "D5")
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 18


def _build_monthly_trend(sheet, variance: VarianceResult) -> None:
    _title(sheet, "月度预算执行趋势", "预算与实际仅比较已经发生的月份")
    monthly = (
        variance.detail.groupby("month", observed=True)[
            ["budget_revenue", "actual_revenue", "budget_operating_profit", "actual_operating_profit"]
        ]
        .sum()
        .reset_index()
    )
    monthly.columns = ["月份", "预算收入", "实际收入", "预算经营利润", "实际经营利润"]
    monthly["月份"] = monthly["月份"].dt.strftime("%Y-%m")
    _write_dataframe(sheet, monthly, "MonthlyTrend", start_row=5)
    for row in range(6, 6 + len(monthly)):
        for column in range(2, 6):
            sheet.cell(row, column).number_format = '¥#,##0;[Red](¥#,##0);-'
    chart = LineChart()
    chart.title = "月度收入预算与实际"
    chart.y_axis.title = "人民币元"
    chart.add_data(Reference(sheet, min_col=2, max_col=3, min_row=5, max_row=11), titles_from_data=True)
    chart.series[0].tx = SeriesLabel(v="预算收入")
    chart.series[1].tx = SeriesLabel(v="实际收入")
    chart.set_categories(Reference(sheet, min_col=1, min_row=6, max_row=11))
    chart.height = 8
    chart.width = 16
    sheet.add_chart(chart, "G5")


def _build_forecast_summary(sheet, forecast: ForecastResult, annual_target: float) -> None:
    _title(sheet, "全年滚动预测情景", f"实际截至 {forecast.actual_through:%Y-%m}｜未来月份按情景假设计算")
    summary = forecast.summary.copy()
    summary["target_operating_profit"] = annual_target
    summary["target_gap"] = summary["operating_profit"] - annual_target
    summary["target_attainment"] = summary["operating_profit"] / annual_target
    summary = summary.rename(
        columns={
            "scenario": "情景",
            "quantity": "数量",
            "revenue": "收入",
            "variable_cost": "变动成本",
            "fixed_expense": "固定费用",
            "gross_profit": "毛利",
            "operating_profit": "经营利润",
            "target_operating_profit": "年度经营利润目标",
            "target_gap": "目标差额",
            "target_attainment": "目标达成率",
        }
    )
    _write_dataframe(sheet, summary, "ForecastSummary", start_row=5)
    for row in range(6, 6 + len(summary)):
        sheet.cell(row, 2).number_format = "#,##0"
        for column in range(3, 10):
            sheet.cell(row, column).number_format = '¥#,##0;[Red](¥#,##0);-'
        sheet.cell(row, 10).number_format = "0.0%"
    chart = BarChart()
    chart.title = "各情景全年经营利润"
    chart.add_data(Reference(sheet, min_col=7, min_row=5, max_row=8), titles_from_data=True)
    chart.series[0].tx = SeriesLabel(v="经营利润")
    chart.set_categories(Reference(sheet, min_col=1, min_row=6, max_row=8))
    chart.height = 8
    chart.width = 15
    sheet.add_chart(chart, "L5")


def _build_data_quality(
    sheet,
    dataset: SyntheticDataset,
    variance: VarianceResult,
    forecast: ForecastResult,
) -> None:
    _title(sheet, "数据质量与模型勾稽", "检查项只用于审阅，不参与业务结果计算")
    checks = [
        ["预算主键重复数", int(dataset.budget.duplicated(["month", "business_unit", "product"]).sum()), 0],
        ["实际主键重复数", int(dataset.actual.duplicated(["month", "business_unit", "product"]).sum()), 0],
        ["利润桥勾稽差异", variance.reconciliation_difference, 0],
        ["预测明细行数", len(forecast.detail), len(dataset.budget) * len(dataset.assumptions)],
        ["实际截止月份", forecast.actual_through.strftime("%Y-%m"), "2026-06"],
    ]
    sheet.append([])
    sheet.append(["检查项", "结果", "预期", "状态"])
    for label, result, expected in checks:
        status = "通过" if (abs(result - expected) < 1e-8 if isinstance(result, (int, float)) else result == expected) else "异常"
        sheet.append([label, result, expected, status])
    _style_table_area(sheet, 5, 5 + len(checks), 4)
    for row in range(6, 6 + len(checks)):
        sheet.cell(row, 2).number_format = "0.00"
        sheet.cell(row, 3).number_format = "0.00"
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 18
    sheet.column_dimensions["D"].width = 12


def _annual_budget_operating_profit(budget: pd.DataFrame) -> float:
    revenue = budget["budget_quantity"] * budget["budget_unit_price"]
    variable_cost = budget["budget_quantity"] * budget["budget_unit_variable_cost"]
    return float((revenue - variable_cost - budget["budget_fixed_expense"]).sum())


def _forecast_export(frame: pd.DataFrame) -> pd.DataFrame:
    return frame.rename(
        columns={
            "scenario": "情景", "month": "月份", "business_unit": "事业部",
            "product": "产品", "data_type": "数据类型", "quantity": "数量",
            "revenue": "收入", "variable_cost": "变动成本", "fixed_expense": "固定费用",
            "gross_profit": "毛利", "operating_profit": "经营利润",
        }
    )


def _budget_input_export(frame: pd.DataFrame) -> pd.DataFrame:
    return frame.rename(
        columns={
            "month": "月份", "business_unit": "事业部", "product": "产品",
            "budget_quantity": "预算数量", "budget_unit_price": "预算单价",
            "budget_unit_variable_cost": "预算单位变动成本",
            "budget_fixed_expense": "预算固定费用",
        }
    )


def _actual_input_export(frame: pd.DataFrame) -> pd.DataFrame:
    return frame.rename(
        columns={
            "month": "月份", "business_unit": "事业部", "product": "产品",
            "actual_quantity": "实际数量", "actual_revenue": "实际收入",
            "actual_variable_cost": "实际变动成本", "actual_fixed_expense": "实际固定费用",
        }
    )


def _assumption_input_export(frame: pd.DataFrame) -> pd.DataFrame:
    return frame.rename(
        columns={
            "scenario": "情景", "quantity_factor": "数量系数",
            "price_factor": "价格系数", "unit_cost_factor": "单位成本系数",
            "fixed_expense_factor": "固定费用系数",
        }
    )


def _variance_export(frame: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "month", "business_unit", "product", "budget_quantity", "actual_quantity",
        "budget_revenue", "actual_revenue", "revenue_variance",
        "budget_operating_profit", "actual_operating_profit", "operating_profit_variance",
        "quantity_profit_impact", "price_profit_impact", "unit_cost_profit_impact",
        "fixed_expense_profit_impact", "reconciliation_difference",
    ]
    labels = [
        "月份", "事业部", "产品", "预算数量", "实际数量", "预算收入", "实际收入",
        "收入差异", "预算经营利润", "实际经营利润", "经营利润差异", "数量与结构合计影响",
        "价格影响", "单位成本影响", "固定费用影响", "勾稽差异",
    ]
    result = frame[columns].copy()
    result.columns = labels
    return result


def _write_dataframe(sheet, frame: pd.DataFrame, table_name: str, start_row: int = 1) -> None:
    headers = list(frame.columns)
    for column, header in enumerate(headers, start=1):
        sheet.cell(start_row, column, header)
    for row_number, row in enumerate(frame.itertuples(index=False, name=None), start=start_row + 1):
        for column, value in enumerate(row, start=1):
            if pd.isna(value):
                value = None
            elif isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            elif isinstance(value, Real) and not isinstance(value, bool):
                value = round(float(value), 2)
            sheet.cell(row_number, column, value)
            header = str(headers[column - 1])
            if value is not None and hasattr(value, "year") and hasattr(value, "month"):
                sheet.cell(row_number, column).number_format = "mmm-yy"
            elif "率" in header or "factor" in header:
                sheet.cell(row_number, column).number_format = "0.0%"
            elif any(
                keyword in header
                for keyword in ("收入", "成本", "费用", "利润", "差异", "影响", "金额")
            ):
                sheet.cell(row_number, column).number_format = '¥#,##0;[Red](¥#,##0);-'
            elif "数量" in header or "quantity" in header:
                sheet.cell(row_number, column).number_format = "#,##0.0"
    end_row = start_row + len(frame)
    end_column = len(headers)
    if len(frame):
        reference = f"A{start_row}:{get_column_letter(end_column)}{end_row}"
        table = Table(displayName=table_name, ref=reference)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)
    _style_table_area(sheet, start_row, end_row, end_column)


def _title(sheet, title: str, subtitle: str) -> None:
    sheet["A2"] = title
    sheet["A2"].font = Font(name="Microsoft YaHei", size=15, bold=True, color=NAVY)
    sheet["A3"] = subtitle
    sheet["A3"].font = Font(name="Microsoft YaHei", size=10, italic=True, color="666666")


def _style_table_area(sheet, header_row: int, end_row: int, end_column: int) -> None:
    for cell in sheet[header_row][:end_column]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(
        min_row=header_row + 1, max_row=end_row, min_col=1, max_col=end_column
    ):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color=BLACK)
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=THIN_GRAY)
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(end_column)}{end_row}"


def _finish_workbook(workbook: Workbook) -> None:
    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        for column_cells in sheet.columns:
            letter = get_column_letter(column_cells[0].column)
            width = max(len(str(cell.value or "")) for cell in column_cells) + 3
            current_width = sheet.column_dimensions[letter].width or 0
            sheet.column_dimensions[letter].width = max(
                current_width, min(max(width, 11), 28)
            )
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and cell.font.name != "Microsoft YaHei":
                    cell.font = Font(
                        name="Microsoft YaHei", size=cell.font.sz or 10,
                        bold=cell.font.bold, italic=cell.font.italic,
                        color=cell.font.color,
                    )
