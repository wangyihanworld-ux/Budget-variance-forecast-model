import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from budget_variance_forecast.reporting import (
    INPUT_SHEETS,
    REPORT_SHEETS,
    generate_demo_artifacts,
)


class ReportingTest(unittest.TestCase):
    def test_demo_generates_two_reopenable_workbooks(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            artifacts = generate_demo_artifacts(temporary)
            self.assertTrue(artifacts.input_path.exists())
            self.assertTrue(artifacts.report_path.exists())

            input_book = load_workbook(artifacts.input_path, read_only=False)
            report_book = load_workbook(artifacts.report_path, read_only=False)
            self.assertEqual(input_book.sheetnames, INPUT_SHEETS)
            self.assertEqual(report_book.sheetnames, REPORT_SHEETS)
            input_book.close()
            report_book.close()

    def test_report_contains_charts_tables_and_audit_checks(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            artifacts = generate_demo_artifacts(temporary)
            workbook = load_workbook(artifacts.report_path, read_only=False)
            self.assertEqual(len(workbook["利润差异桥"]._charts), 1)
            self.assertEqual(len(workbook["月度趋势"]._charts), 1)
            self.assertEqual(len(workbook["滚动预测"]._charts), 1)
            self.assertGreater(len(workbook["预测明细"].tables), 0)
            statuses = [workbook["数据质量"].cell(row, 4).value for row in range(6, 11)]
            self.assertEqual(statuses, ["通过"] * 5)
            summary = workbook["管理摘要"]
            self.assertAlmostEqual(summary["B7"].value, 17_584.0, places=2)
            self.assertAlmostEqual(summary["B9"].value, -10_724.36, places=2)
            self.assertAlmostEqual(summary["B13"].value, -10_724.36, places=2)
            workbook.close()

    def test_workbooks_are_visibly_labeled_as_synthetic(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            artifacts = generate_demo_artifacts(temporary)
            workbook = load_workbook(artifacts.input_path, read_only=True)
            messages = [row[0] for row in workbook["演示说明"].iter_rows(min_row=2, values_only=True)]
            self.assertTrue(any("合成数据" in str(message) for message in messages))
            workbook.close()


if __name__ == "__main__":
    unittest.main()
